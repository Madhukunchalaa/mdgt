import os
import openpyxl
from django.core.management.base import BaseCommand
from Employee.models import Employee
from MaterialType.models import MaterialType
from matgroups.models import MatGroup
from matg_attributes.models import MatgAttributeItem
from matg_attributes import catalog
from itemmaster.models import ItemMaster
from itemmaster.views import format_short_name, format_long_name


DEFAULT_EXCEL_PATH = r"C:\Users\IE-Admin\Desktop\Sample Material group code_piping_28092026.xlsx"

CHAR_SPECS = [
    ("PART TYPE", 10, 8, 9),
    ("SUB TYPE", 20, 10, 11),
    ("SIZE", 30, 12, 13),
    ("SCHEDULE", 40, 14, 15),
    ("RATING", 50, 16, 17),
    ("MOC", 60, 18, 19),
    ("STANDARD", 70, 20, 21),
]


class Command(BaseCommand):
    help = "Import Piping Material Groups, Characteristics with CMT codes, and Items from Excel"

    def add_arguments(self, parser):
        parser.add_argument("--file", default=DEFAULT_EXCEL_PATH, help="Path to the Excel file")
        parser.add_argument("--dry-run", action="store_true", help="Parse without saving to database")

    def handle(self, *args, **options):
        file_path = options["file"]
        dry_run = options["dry_run"]

        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File not found: {file_path}"))
            return

        self.stdout.write(self.style.MIGRATE_HEADING(f"Loading Excel file: {file_path}"))
        wb = openpyxl.load_workbook(file_path, data_only=True)

        if "Piping" not in wb.sheetnames:
            self.stderr.write(self.style.ERROR("Sheet 'Piping' not found in workbook!"))
            return

        emp = Employee.objects.filter(email="admin@meil.com").first() or Employee.objects.first()
        mat_type = MaterialType.objects.filter(mat_type_code="ZCON").first() or MaterialType.objects.first()

        ws_piping = wb["Piping"]

        # -------------------------------------------------------------
        # 1. Parse Sheet1 for master values and CMT codes (if present)
        # -------------------------------------------------------------
        master_chars = {name: {"values": [], "codes": {}} for name, _, _, _ in CHAR_SPECS}
        if "Sheet1" in wb.sheetnames:
            ws_s1 = wb["Sheet1"]
            # Sheet1 layout: cols 2-3 (Part Type), 4-5 (Sub Type), 6-7 (Size), 8-9 (Schedule), 10-11 (Rating), 12-13 (MOC), 14-15 (Standard)
            sheet1_map = [
                ("PART TYPE", 2, 3),
                ("SUB TYPE", 4, 5),
                ("SIZE", 6, 7),
                ("SCHEDULE", 8, 9),
                ("RATING", 10, 11),
                ("MOC", 12, 13),
                ("STANDARD", 14, 15),
            ]
            for r in range(3, ws_s1.max_row + 1):
                for c_name, code_col, val_col in sheet1_map:
                    raw_c = ws_s1.cell(r, code_col).value
                    raw_v = ws_s1.cell(r, val_col).value
                    if raw_v is not None and str(raw_v).strip():
                        v_str = str(raw_v).strip()
                        c_str = str(raw_c).strip() if raw_c is not None else ""
                        if len(c_str) == 1 and c_str.isdigit():
                            c_str = "0" + c_str
                        if v_str not in master_chars[c_name]["values"]:
                            master_chars[c_name]["values"].append(v_str)
                        if c_str and c_str != "00":
                            master_chars[c_name]["codes"][v_str] = c_str

        # -------------------------------------------------------------
        # 2. Parse Piping tab for Material Groups and Items
        # -------------------------------------------------------------
        groups_dict = {}
        rows_data = []

        for r in range(3, ws_piping.max_row + 1):
            mgrp_code = ws_piping.cell(r, 24).value
            raw_code = ws_piping.cell(r, 22).value
            desc = ws_piping.cell(r, 23).value
            mgrp_desc = ws_piping.cell(r, 25).value

            if not mgrp_code and not desc and not raw_code:
                continue
            if not mgrp_code:
                continue

            mgrp_code = str(mgrp_code).strip()
            if mgrp_code not in groups_dict:
                groups_dict[mgrp_code] = {
                    "desc": str(mgrp_desc or mgrp_code).strip(),
                    "chars": {
                        name: {
                            "seq": seq,
                            "values": list(master_chars[name]["values"]),
                            "codes": dict(master_chars[name]["codes"]),
                        }
                        for name, seq, _, _ in CHAR_SPECS
                    },
                }

            row_attrs = {}
            for c_name, seq, c_code_col, c_val_col in CHAR_SPECS:
                raw_c = ws_piping.cell(r, c_code_col).value
                raw_v = ws_piping.cell(r, c_val_col).value

                c_str = str(raw_c).strip() if raw_c is not None else ""
                if len(c_str) == 1 and c_str.isdigit():
                    c_str = "0" + c_str

                v_str = str(raw_v).strip() if raw_v is not None else ""
                if v_str:
                    row_attrs[c_name] = v_str
                    char_bucket = groups_dict[mgrp_code]["chars"][c_name]
                    if v_str not in char_bucket["values"]:
                        char_bucket["values"].append(v_str)
                    if c_str and c_str != "00":
                        char_bucket["codes"][v_str] = c_str

            clean_code = str(raw_code).strip() if raw_code is not None else ""
            if clean_code.isdigit() and len(clean_code) % 2 != 0:
                clean_code = "0" + clean_code  # restore leading zero stripped by Excel

            rows_data.append({
                "row": r,
                "mgrp_code": mgrp_code,
                "excel_code": clean_code,
                "desc": str(desc or "").strip(),
                "attrs": row_attrs,
            })

        self.stdout.write(f"Found {len(groups_dict)} Material Groups and {len(rows_data)} Item rows.")

        if dry_run:
            self.stdout.write(self.style.WARNING("DRY RUN: No database changes were made."))
            return

        # -------------------------------------------------------------
        # 3. Create/Update Material Groups & Characteristics in DB
        # -------------------------------------------------------------
        created_groups = 0
        created_chars = 0
        for mgrp_code, g_data in groups_dict.items():
            shortname = g_data["desc"]
            if len(shortname) > 150:
                shortname = shortname[:150]

            grp, g_created = MatGroup.objects.update_or_create(
                mgrp_code=mgrp_code,
                defaults={
                    "mgrp_shortname": shortname,
                    "mgrp_longname": g_data["desc"][:150],
                    "search_type": "Materials",
                    "createdby": emp,
                    "updatedby": emp,
                    "is_deleted": False,
                },
            )
            if g_created:
                created_groups += 1

            for c_name, c_info in g_data["chars"].items():
                _, c_created = MatgAttributeItem.objects.update_or_create(
                    mgrp_code=grp,
                    attribute_name=c_name,
                    defaults={
                        "possible_values": c_info["values"],
                        "cmt_codes": c_info["codes"],
                        "print_priority": c_info["seq"],
                        "createdby": emp,
                        "updatedby": emp,
                        "is_deleted": False,
                    },
                )
                if c_created:
                    created_chars += 1

        self.stdout.write(self.style.SUCCESS(
            f"Groups processed: {len(groups_dict)} ({created_groups} new); "
            f"Characteristics configured: {len(groups_dict) * len(CHAR_SPECS)} ({created_chars} new)"
        ))

        # -------------------------------------------------------------
        # 4. Create/Update Items in DB
        # -------------------------------------------------------------
        created_items = 0
        updated_items = 0

        for row in rows_data:
            grp = MatGroup.objects.get(mgrp_code=row["mgrp_code"])
            attrs = catalog.order_attributes(grp, row["attrs"])
            built_code = catalog.build_catalog_number(grp, attrs)

            # Prefer built catalog number if clean, fallback to normalized excel_code
            final_catalog_number = built_code or (row["excel_code"] if row["excel_code"] != "x" else None)

            base_sap_name = row["attrs"].get("PART TYPE") or row["desc"] or grp.mgrp_shortname
            short_name = format_short_name(base_sap_name, attrs)
            if not short_name or len(short_name) < 3:
                short_name = row["desc"] or base_sap_name

            long_name = format_long_name(base_sap_name, grp.mgrp_code, grp.mgrp_longname, attrs)

            search_text_parts = [row["desc"], final_catalog_number or ""]
            search_text = " ".join([p for p in search_text_parts if p])

            item, i_created = ItemMaster.objects.update_or_create(
                mgrp_code=grp,
                short_name=short_name[:100],
                defaults={
                    "mat_type_code": mat_type,
                    "sap_name": base_sap_name[:150],
                    "long_name": long_name[:250],
                    "mgrp_long_name": grp.mgrp_longname,
                    "search_text": search_text[:300],
                    "attributes": attrs,
                    "catalog_number": final_catalog_number,
                    "createdby": emp,
                    "updatedby": emp,
                    "is_deleted": False,
                },
            )
            if i_created:
                created_items += 1
            else:
                updated_items += 1

        self.stdout.write(self.style.SUCCESS(
            f"Successfully imported items: {created_items} created, {updated_items} updated."
        ))
